# import os
# import pandas as pd
# import matplotlib.pyplot as plt
# from openpyxl import Workbook
# from openpyxl.drawing.image import Image as XLImage
# from openpyxl.styles import Font
# import io
# import tempfile
#
# # Fix OpenMP runtime conflict (common on Windows)
# os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
#
# from ultralytics import YOLO
#
# #  Train the model
# print("Starting training...")
# model = YOLO('yolov8n.pt')
# results = model.train(
#     data='C:/Users/troll/PycharmProjects/PythonProject5/victim_detect_alt dataset/archive/C2A_Dataset/data.yaml',
#     epochs=100,
#     imgsz=640,
#     batch=16,
#     device='cpu',
#     name='train5',
#     save=True,  # Ensures model is saved automatically
#     plots=True  # Generates default YOLO plots as backup
# )
#
# # YOLO returns the exact save directory path
# run_dir = str(results.save_dir)
# csv_path = os.path.join(run_dir, 'results.csv')
# weights_dir = os.path.join(run_dir, 'weights')
#
# if not os.path.exists(csv_path):
#     raise FileNotFoundError("Training failed or results.csv not found. Check console for errors.")
#
# df = pd.read_csv(csv_path)
#
# #  Create Excel Workbook
# wb = Workbook()
#
# #Sheet 1: Raw Training Numbers
# ws_data = wb.active
# ws_data.title = "Training Metrics"
# for c_idx, col in enumerate(df.columns, 1):
#     ws_data.cell(row=1, column=c_idx, value=col)
# for r_idx, row in enumerate(df.values, 2):
#     for c_idx, val in enumerate(row, 1):
#         ws_data.cell(row=r_idx, column=c_idx, value=val)
#
# #Sheet 2: Graphs
# ws_graphs = wb.create_sheet("Training Graphs")
#
# # Define metric pairs to plot (Train vs Val or key metrics)
# plot_configs = [
#     ("train/box_loss", "val/box_loss", "Box Loss (Train vs Validation)"),
#     ("train/cls_loss", "val/cls_loss", "Class Loss (Train vs Validation)"),
#     ("metrics/mAP50(B)", "metrics/mAP50-95(B)", "mAP Metrics"),
#     ("lr/pg0", "lr/pg1", "Learning Rates")
# ]
#
# row_offset = 2
# for col1, col2, title in plot_configs:
#     if col1 in df.columns and col2 in df.columns:
#         plt.figure(figsize=(8, 4.5))
#         plt.plot(df['epoch'], df[col1], label=f'Train', linewidth=2)
#         plt.plot(df['epoch'], df[col2], label=f'Validation', linewidth=2)
#         plt.title(title, fontsize=12, fontweight='bold')
#         plt.xlabel('Epoch')
#         plt.ylabel('Value')
#         plt.legend()
#         plt.grid(True, alpha=0.3)
#         plt.tight_layout()
#
#         # Save plot to memory buffer
#         buf = io.BytesIO()
#         plt.savefig(buf, format='png', dpi=150)
#         buf.seek(0)
#         plt.close()
#
#         # Create temporary image file for openpyxl
#         with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp:
#             tmp.write(buf.read())
#             tmp_path = tmp.name
#
#         # Embed in Excel
#         ws_graphs.cell(row=row_offset, column=1, value=title).font = Font(bold=True, size=12)
#         row_offset += 1
#
#         img_obj = XLImage(tmp_path)
#         img_obj.width = 480
#         img_obj.height = 290
#         ws_graphs.add_image(img_obj, f"A{row_offset}")
#         row_offset += 20  # Space between graphs
#
#         # Cleanup temp file
#         os.remove(tmp_path)
#     else:
#         print(f"Skipping plot: {title} (columns not found in results.csv)")
#
# #  Save Excel Report
# excel_path = os.path.join(run_dir, 'training_progress_report.xlsx')
# wb.save(excel_path)
#
# print(f"\n Training Complete!")
# print(f" Trained model saved in: {weights_dir}")
# print(f"   ├─ best.pt (highest validation mAP)")
# print(f"   └─ last.pt (final epoch weights)")
# print(f" Excel report with numbers & graphs: {excel_path}")

import os
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment
import shutil

os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"

from ultralytics import YOLO

# Configuration
DATA_YAML = 'C:/Users/troll/PycharmProjects/PythonProject5/victim_detect_alt dataset/archive/C2A_Dataset/data.yaml'
MODEL_PATH = 'yolov8n.pt'
EPOCHS = 20
IMGSZ = 640
BATCH = 16
DEVICE = 'cpu'
RUN_NAME = 'train59'

print("Starting training...")
model = YOLO(MODEL_PATH)
results = model.train(
    data=DATA_YAML,
    epochs=EPOCHS,
    imgsz=IMGSZ,
    batch=BATCH,
    device=DEVICE,
    name=RUN_NAME,
    save=True,
    plots=True
)

run_dir = str(results.save_dir)
csv_path = os.path.join(run_dir, 'results.csv')
weights_dir = os.path.join(run_dir, 'weights')
best_model_path = os.path.join(weights_dir, 'best.pt')

if not os.path.exists(csv_path):
    raise FileNotFoundError("Training failed or results.csv not found.")

df = pd.read_csv(csv_path)

# Run testing on the test set using the best model
print("Running testing on test set...")
test_results = model.val(
    data=DATA_YAML,
    split='test',
    imgsz=IMGSZ,
    batch=BATCH,
    device=DEVICE,
    save_json=True,
    name=f'{RUN_NAME}_test'
)

# Extract test metrics
test_metrics = {
    'test_precision': float(test_results.results_dict.get('metrics/precision(B)', 0)),
    'test_recall': float(test_results.results_dict.get('metrics/recall(B)', 0)),
    'test_mAP50': float(test_results.results_dict.get('metrics/mAP50(B)', 0)),
    'test_mAP50_95': float(test_results.results_dict.get('metrics/mAP50-95(B)', 0))
}

print(f"Test Results - Precision: {test_metrics['test_precision']:.4f}, "
      f"Recall: {test_metrics['test_recall']:.4f}, "
      f"mAP50: {test_metrics['test_mAP50']:.4f}, "
      f"mAP50-95: {test_metrics['test_mAP50_95']:.4f}")

# Create Excel Workbook
wb = Workbook()

# Sheet 1: Training Metrics (Raw Data)
ws_data = wb.active
ws_data.title = "Training_Metrics"
headers = list(df.columns)
for c_idx, col in enumerate(headers, 1):
    cell = ws_data.cell(row=1, column=c_idx, value=col)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')
for r_idx, row in enumerate(df.values, 2):
    for c_idx, val in enumerate(row, 1):
        ws_data.cell(row=r_idx, column=c_idx, value=val)

# Sheet 2: Test Results Summary
ws_test = wb.create_sheet("Test_Results")
ws_test.cell(row=1, column=1, value="Metric").font = Font(bold=True)
ws_test.cell(row=1, column=2, value="Value").font = Font(bold=True)
for idx, (metric, value) in enumerate(test_metrics.items(), 2):
    ws_test.cell(row=idx, column=1, value=metric.replace('test_', ''))
    ws_test.cell(row=idx, column=2, value=value)

# Sheet 3: Graphs
ws_graphs = wb.create_sheet("Training_Graphs")

plot_configs = [
    ("train/box_loss", "val/box_loss", "Box_Loss_Train_vs_Val"),
    ("train/cls_loss", "val/cls_loss", "Class_Loss_Train_vs_Val"),
    ("metrics/mAP50(B)", "metrics/mAP50-95(B)", "mAP_Metrics"),
    ("lr/pg0", "lr/pg1", "Learning_Rates")
]

plot_folder = os.path.join(run_dir, 'excel_plots')
os.makedirs(plot_folder, exist_ok=True)

row_offset = 2
for col1, col2, title in plot_configs:
    if col1 in df.columns and col2 in df.columns:
        plt.figure(figsize=(8, 4.5))
        plt.plot(df['epoch'], df[col1], label='Train', linewidth=2)
        plt.plot(df['epoch'], df[col2], label='Validation', linewidth=2)
        plt.title(title.replace('_', ' '), fontsize=12, fontweight='bold')
        plt.xlabel('Epoch')
        plt.ylabel('Value')
        plt.legend()
        plt.grid(True, alpha=0.3)
        plt.tight_layout()

        plot_filename = f"{title}.png"
        plot_path = os.path.join(plot_folder, plot_filename)
        plt.savefig(plot_path, dpi=150)
        plt.close()

        ws_graphs.cell(row=row_offset, column=1, value=title.replace('_', ' ')).font = Font(bold=True, size=11)
        row_offset += 1

        img_obj = XLImage(plot_path)
        img_obj.width = 480
        img_obj.height = 290
        ws_graphs.add_image(img_obj, f"A{row_offset}")
        row_offset += 20
    else:
        print(f"Skipping plot: {title} (columns not found in results.csv)")

# Save Excel file
excel_path = os.path.join(run_dir, 'training_progress_report.xlsx')
wb.save(excel_path)

# Cleanup temporary plot files
if os.path.exists(plot_folder):
    shutil.rmtree(plot_folder)

print(f"Training complete.")
print(f"Model saved: {best_model_path}")
print(f"Excel report: {excel_path}")